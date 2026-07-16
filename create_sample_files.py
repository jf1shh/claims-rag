import os
import docx
import pandas as pd

# Create output folder
os.makedirs("sample_guidelines", exist_ok=True)

# 1. Generate DOCX: Auto Insurance Claims Policy Guidelines
def create_auto_claims_docx():
    doc = docx.Document()
    doc.add_heading("Standard Operating Guidelines: Auto Insurance Claims", level=1)
    
    doc.add_heading("1. Comprehensive vs. Collision Coverage", level=2)
    doc.add_paragraph(
        "Collision Coverage applies to damage to the policyholder's vehicle resulting from a collision with "
        "another vehicle or a stationary object (e.g., trees, guardrails, barriers). Comprehensive Coverage "
        "applies to damage caused by events other than collision, including theft, vandalism, fire, contact with "
        "animals (e.g., hitting a deer), windstorms, hail, or falling objects (e.g., tree branches)."
    )
    
    doc.add_heading("2. Claims Processing Requirements", level=2)
    doc.add_paragraph(
        "To process any physical damage auto claim, the claims handler must verify and compile the following evidence:"
    )
    doc.add_paragraph("- Official Police Report: Required for all multi-vehicle accidents, thefts, and hit-and-runs.")
    doc.add_paragraph("- Clear Damage Photographs: High-resolution images of all four sides of the vehicle, the odometer, the VIN plate, and close-ups of the direct impact zones.")
    doc.add_paragraph("- Detailed Repair Estimate: A line-item breakdown from an approved network body shop, or a licensed independent adjuster report.")
    doc.add_paragraph("- Dashcam or Eyewitness Statements: Highly recommended for liability disputes.")

    doc.add_heading("3. Glass & Windshield Claims Policy", level=2)
    doc.add_paragraph(
        "Windshield and safety glass replacement is processed under Comprehensive coverage. Under standard guidelines, "
        "if the windshield can be repaired rather than replaced (e.g., chips smaller than a quarter), the deductible "
        "is fully waived. If full replacement is required, the standard Comprehensive deductible applies unless "
        "the policyholder has purchased a Zero-Deductible Glass endorsement."
    )
    
    doc.add_heading("4. Rental Car Reimbursement Rules", level=2)
    doc.add_paragraph(
        "Rental reimbursement coverage (Transportation Expenses) is optional and only applies if the vehicle is "
        "rendered inoperable due to a covered loss. Under standard plans, reimbursement is capped at $30 per day "
        "for a maximum of 30 days ($900 total limit). Premium plans extend this limit to $50 per day for up to "
        "30 days ($1,500 total limit). Reimbursement is only provided during the actual period of active repairs."
    )
    
    doc.save("sample_guidelines/Auto_Claims_Guidelines.docx")
    print("Created: sample_guidelines/Auto_Claims_Guidelines.docx")


# 2. Generate XLSX: Auto Deductibles and Limits Schedule
def create_auto_limits_xlsx():
    data = {
        "Policy Plan": [
            "Standard Auto (Plan A)", 
            "Standard Auto (Plan A)", 
            "Premium Auto (Plan B)", 
            "Premium Auto (Plan B)", 
            "Basic Auto (Plan C)"
        ],
        "Coverage Category": [
            "Collision Damage", 
            "Comprehensive Perils", 
            "Collision Damage", 
            "Comprehensive Perils", 
            "Collision Only"
        ],
        "Coverage Limit ($)": [
            "Actual Cash Value (ACV)", 
            "Actual Cash Value (ACV)", 
            "Actual Cash Value (ACV)", 
            "Actual Cash Value (ACV)", 
            "Actual Cash Value (ACV)"
        ],
        "Standard Deductible ($)": [
            1000, 
            500, 
            500, 
            250, 
            2000
        ],
        "Towing & Roadside Assistance": [
            "Capped at $75/event", 
            "Capped at $75/event", 
            "Full Coverage (Unlimited)", 
            "Full Coverage (Unlimited)", 
            "Not Covered"
        ],
        "Rental Reimbursement ($/Day)": [
            30, 
            30, 
            50, 
            50, 
            0
        ]
    }
    df = pd.DataFrame(data)
    df.to_excel("sample_guidelines/Auto_Deductibles_And_Limits.xlsx", index=False)
    print("Created: sample_guidelines/Auto_Deductibles_And_Limits.xlsx")


# 3. Generate General Exclusions txt file for Auto
def create_auto_exclusions_txt():
    content = """=========================================
AUTO INSURANCE POLICY - PERIL EXCLUSIONS
=========================================

Section 12: Exclusions from Physical Damage Coverage

The insurer will not pay for loss or damage caused by, resulting from, or consisting of the following:

12.1. Intentional Acts: Damage caused by or at the direction of the policyholder, or any relative residing in the household, with the intent to cause damage or defraud the insurer.

12.2. Under the Influence (DUI/DWI): Damage arising from any collision where the operator of the insured vehicle is cited for and subsequently convicted of driving under the influence of alcohol, narcotics, or illegal substances at the time of the occurrence.

12.3. Racing and Stunting: Loss occurring while the insured vehicle is being operated in any pre-arranged or organized racing contest, speed test, demolition derby, or performance stunting event.

12.4. Commercial Delivery / Ride-Sharing: Damage occurring while the vehicle is being used for commercial purposes, including cargo transport, public transportation, package delivery, or active ride-sharing (e.g., Uber, Lyft) unless a specific commercial business endorsement has been added to the policy.

12.5. Wear and Tear / Mechanical Breakdown: Any damage resulting from road wear, rust, corrosion, manufacturer defects, or electrical/mechanical failures not directly resulting from an external covered accident or peril.
"""
    with open("sample_guidelines/Auto_Policy_Exclusions.txt", "w", encoding="utf-8") as f:
        f.write(content)
    print("Created: sample_guidelines/Auto_Policy_Exclusions.txt")

if __name__ == "__main__":
    create_auto_claims_docx()
    create_auto_limits_xlsx()
    create_auto_exclusions_txt()
